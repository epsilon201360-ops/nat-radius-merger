"""Export functionality for merged data."""

import csv
from typing import List

from app.models import MergedRecord
from app.utils import format_date

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class DataExporter:
    """Export merged records to CSV or Excel."""

    @staticmethod
    def export_to_csv(merged_records: List[MergedRecord], output_path: str) -> bool:
        """Export merged records as CSV."""
        try:
            with open(output_path, "w", newline="", encoding="utf-8") as file:
                fieldnames = [
                    "NAT Time",
                    "MSISDN",
                    "IP Private",
                    "NAT Type",
                    "Event Type",
                    "Original Source IP",
                    "Translated IP",
                    "Destination IP",
                    "Device Name",
                    "Protocol",
                    "RADIUS State",
                ]
                writer = csv.DictWriter(file, fieldnames=fieldnames)
                writer.writeheader()

                for record in merged_records:
                    writer.writerow(
                        {
                            "NAT Time": format_date(record.nat_time),
                            "MSISDN": record.msisdn,
                            "IP Private": record.ip_private,
                            "NAT Type": record.nat_type,
                            "Event Type": record.event_type,
                            "Original Source IP": record.original_source_ip,
                            "Translated IP": record.translated_ip,
                            "Destination IP": record.destination_ip,
                            "Device Name": record.device_name,
                            "Protocol": record.protocol,
                            "RADIUS State": "Connection" if record.radius_state == 1 else "Disconnection",
                        }
                    )

            return True
        except Exception as exc:
            print(f"Error exporting to CSV: {exc}")
            return False

    @staticmethod
    def export_to_excel(merged_records: List[MergedRecord], output_path: str) -> bool:
        """Export merged records as Excel."""
        if not OPENPYXL_AVAILABLE:
            print("openpyxl not available. Install it with: pip install openpyxl")
            return False

        try:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "NAT-RADIUS Merge"

            headers = [
                "NAT Time",
                "MSISDN",
                "IP Private",
                "NAT Type",
                "Event Type",
                "Original Source IP",
                "Translated IP",
                "Destination IP",
                "Device Name",
                "Protocol",
                "RADIUS State",
            ]

            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            header_font = Font(color="FFFFFF", bold=True)

            for column_index, header in enumerate(headers, 1):
                cell = sheet.cell(row=1, column=column_index)
                cell.value = header
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center")

            for row_index, record in enumerate(merged_records, 2):
                sheet.cell(row=row_index, column=1).value = format_date(record.nat_time)
                sheet.cell(row=row_index, column=2).value = record.msisdn
                sheet.cell(row=row_index, column=3).value = record.ip_private
                sheet.cell(row=row_index, column=4).value = record.nat_type
                sheet.cell(row=row_index, column=5).value = record.event_type
                sheet.cell(row=row_index, column=6).value = record.original_source_ip
                sheet.cell(row=row_index, column=7).value = record.translated_ip
                sheet.cell(row=row_index, column=8).value = record.destination_ip
                sheet.cell(row=row_index, column=9).value = record.device_name
                sheet.cell(row=row_index, column=10).value = record.protocol
                sheet.cell(row=row_index, column=11).value = "Connection" if record.radius_state == 1 else "Disconnection"

            for column in range(1, len(headers) + 1):
                sheet.column_dimensions[chr(64 + column)].width = 20

            workbook.save(output_path)
            return True
        except Exception as exc:
            print(f"Error exporting to Excel: {exc}")
            return False
